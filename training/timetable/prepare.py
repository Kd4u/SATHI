#!/usr/bin/env python3
"""Prepare PRIVATE reviewed cell-role features from date/day timetable workbooks.

The course-label file must be a reviewed JSON list of canonical subject names.
Unknown non-event text aborts preparation for review; no model labels itself.
Requires openpyxl. Writes features plus private provenance to the output folder.
"""
import argparse
import collections
import datetime
import json
import re
from pathlib import Path
import openpyxl
from train import FEATURES, EVENT_TOKENS


def professor(text):
    return bool(re.match(r"(?i)^\s*(professor|prof\.?|dr\.?|mr\.?|mrs\.?|ms\.?)\s+", text))


def time_only(text):
    return bool(re.fullmatch(r"\s*\d{1,2}[:.]\d{2}\s*[-–—]\s*\d{1,2}[:.]\d{2}\s*", text))


def canonical(text):
    text = re.sub(r"\s+[-–—]\s*\d{1,3}\s*$", "", text)
    text = re.sub(r"(?i)\(?\s*\bsec(?:tion)?\.?\s*[-:]?\s*[A-Z]\b\s*\)?", "", text)
    return " ".join(text.split())


def features(text, faculty, repetitions, known):
    return [int(professor(faculty)), int(known), int(repetitions > 1),
            int(bool(re.match(r"(?i)^TBD\d*$", faculty))), int("GUEST" in faculty.upper()),
            int(any(token in text.upper() for token in EVENT_TOKENS)), int(time_only(text)),
            int(text.strip() in [".", "-", "TBD", ""]), min(len(text.split()), 12) / 12,
            int(bool(re.search(r"\s[-–—]\s*\d+\s*$", text))),
            int(bool(re.search(r"(?i)\bsec(?:tion)?\.?\s*[A-Z]\b", text)))]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--class-labels", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("workbooks", nargs="+", type=Path)
    args = parser.parse_args()
    courses = set(json.loads(args.class_labels.read_text()))
    examples, provenance = [], []
    for fold, path in enumerate(args.workbooks):
        for sheet in openpyxl.load_workbook(path, data_only=True):
            raw = []
            for index, row in enumerate(sheet.values):
                if not isinstance(row[0], datetime.datetime):
                    continue
                for column, value in enumerate(row[2:], 3):
                    if not value:
                        continue
                    entries = [line.strip() for line in str(value).splitlines() if line.strip()]
                    entries = [entry for entry in entries if not time_only(entry)]
                    faculty = [line.strip() for line in str(sheet.cell(index + 2, column).value or "").splitlines() if line.strip()]
                    for entry_index, text in enumerate(entries):
                        teacher = faculty[entry_index] if len(faculty) == len(entries) else faculty[0] if len(faculty) == 1 else ""
                        name = canonical(text)
                        label = int(name in courses)
                        if not label and not (any(token in text.upper() for token in EVENT_TOKENS) or text in [".", "-"]):
                            raise ValueError(f"Review label: {path.name} {sheet.title}!{index + 1}: {text}")
                        source = f"{sheet.title}!{openpyxl.utils.get_column_letter(column)}{index + 1}"
                        raw.append((text, teacher, name, label, source))
            counts = collections.Counter(row[2] for row in raw)
            known = {row[2] for row in raw if professor(row[1])}
            for text, teacher, name, label, source in raw:
                examples.append(dict(fold=fold, x=features(text, teacher, counts[name], name in known), y=label))
                provenance.append(dict(workbook=path.name, source=source, text=text, label="class" if label else "event"))
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / "examples.json").write_text(json.dumps(examples))
    (args.output / "provenance-private.json").write_text(json.dumps(provenance))
    print(f"Prepared {len(examples)} reviewed feature rows; features: {FEATURES}")

if __name__ == "__main__":
    main()
